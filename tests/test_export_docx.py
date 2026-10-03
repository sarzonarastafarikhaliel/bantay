"""The .docx export fills the paper BLOTTER FORM from the stored narrative, and
shows non-admins the same redacted names the record page does."""
import io

import pytest
from docx import Document

from bantay import create_app, db
from bantay import narrative as narrative_tpl
from bantay.models import IncidentRecord, User


@pytest.fixture
def app(tmp_path):
    flask_app = create_app(overrides={
        "SQLALCHEMY_DATABASE_URI": f"sqlite:///{tmp_path / 'test.db'}",
        "TESTING": True,
        "MODEL_DIR": str(tmp_path / "models"),
    })
    with flask_app.app_context():
        for name, role in (("staff", "staff"), ("boss", "admin")):
            u = User(username=name, role=role)
            u.set_password("pw")
            db.session.add(u)
        slots = dict.fromkeys(narrative_tpl.SLOTS, "")
        slots.update(blotter_no="2026-014", reporting_party="Juan Dela Cruz",
                     complaint="Pagnanakaw ng manok", date="2026-01-03",
                     incident_summary="Nagsadya si Juan upang ireklamo ang manok.")
        db.session.add(IncidentRecord(narrative=narrative_tpl.render(slots), date="2026-01-03"))
        db.session.add(IncidentRecord(narrative="Free-text entry, no form layout.",
                                      date="2026-02-01", location_purok="Purok 5"))
        db.session.commit()
    yield flask_app


def _login(app, name):
    c = app.test_client()
    c.post("/login", data={"username": name, "password": "pw"})
    return c


def _text(resp):
    assert resp.status_code == 200
    assert resp.mimetype.endswith("wordprocessingml.document")
    doc = Document(io.BytesIO(resp.data))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        parts += [c.text for row in t.rows for c in row.cells]
    return "\n".join(parts)


def test_admin_sees_form_filled(app):
    text = _text(_login(app, "boss").get("/records/1/export.docx"))
    assert "BLOTTER FORM" in text and "2026-014" in text
    assert "Juan Dela Cruz" in text and "Pagnanakaw ng manok" in text


def test_staff_gets_names_redacted(app):
    text = _text(_login(app, "staff").get("/records/1/export.docx"))
    assert "Juan Dela Cruz" not in text and "[REDACTED]" in text


def test_bulk_follows_filters_and_free_text_fallback(app):
    c = _login(app, "boss")
    text = _text(c.get("/records/export.docx?location=Purok 5"))
    assert "Free-text entry" in text and "Purok 5" in text
    assert "Juan Dela Cruz" not in text
    assert text.count("BLOTTER FORM") == 1
    assert _text(c.get("/records/export.docx")).count("BLOTTER FORM") == 2
